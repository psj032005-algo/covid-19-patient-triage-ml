"""Builds report/Slides.pptx from results/summary.json.  python src/build_slides.py"""
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

root = Path(__file__).resolve().parent.parent
S = json.load(open(root / "results/summary.json")); I, D = S["icu"], S["death"]
X2 = json.load(open(root / "results/improved_summary.json"))
FIG = root / "results/figures"
NAVY, INK, MUTE, ACC, RED = RGBColor(0x1F, 0x3B, 0x6F), RGBColor(0x1F, 0x29, 0x37), RGBColor(0x6B, 0x72, 0x80), RGBColor(0x1F, 0x5F, 0xA8), RGBColor(0xB9, 0x1C, 0x1C)
P = Presentation(); P.slide_width, P.slide_height = Inches(13.333), Inches(7.5)
blank = P.slide_layouts[6]

def tx(sl, x, y, w, h, text, size=18, bold=False, color=INK, align=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, l in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = l; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Calibri"
        p.space_after = Pt(6)
        if align: p.alignment = align
    return tb

def slide(title, sub=None):
    sl = P.slides.add_slide(blank)
    tx(sl, .6, .35, 12.1, .8, title, 30, True, NAVY)
    if sub: tx(sl, .6, 1.05, 12.1, .5, sub, 16, False, MUTE)
    return sl

def bullets(sl, items, x=.6, y=1.7, w=12.1, size=20, h=5):
    tx(sl, x, y, w, h, ["•  " + i for i in items], size)

def img(sl, name, x, y, w=None, h=None):
    sl.shapes.add_picture(str(FIG / name), Inches(x), Inches(y), width=Inches(w) if w else None, height=Inches(h) if h else None)

def box(sl, x, y, w, h, text, fill=ACC, size=15):
    b = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = fill; b.line.fill.background()
    tf = b.text_frame; tf.word_wrap = True; tf.text = text
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs: r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255); r.font.bold = True; r.font.name = "Calibri"

# 1 title
sl = P.slides.add_slide(blank)
tx(sl, .8, 2.2, 11.7, 1.6, "CT-based Patient Triage of COVID-19:\nPredicting ICU Admission and Death", 40, True, NAVY)
tx(sl, .8, 4.2, 11.7, .6, "Problem Statement 12 - adapted to a public lab-data cohort", 22, False, MUTE)
tx(sl, .8, 5.2, 11.7, 1, ["UE24CS352A Machine Learning - Mini-Project - Group 23", "Prajwal S Jyanopanthar (PES1UG24CS594)  |  Prakash N Heroor (PES1UG24CS595)"], 18)

sl = slide("Problem and source paper", "Zhan & Li, CS229 Spring 2020, Stanford")
bullets(sl, ["Goal: at admission, flag COVID-19 patients likely to need ICU, ventilation, or to die - so limited resources go to those at highest risk",
             "Paper: CT radiomics + clinical + lab data; 1662 patients to develop models, 700 unseen patients to validate",
             "Paper's method: SMOTEENN for rare outcomes; LASSO / PCA / FPR feature selection; LR, RF, SVM, MLP, LightGBM; bootstrap comparison of data types",
             "Paper's best: AUROC 0.91 (ICU), 0.92 (ventilation), 0.83 (death) using radiomics + clinical + lab"])

sl = slide("Our adaptation: what we could and could not reproduce")
bullets(sl, ["The paper's data (Jinling Hospital) is private; the authors' GitHub repository returns 404 - original data cannot be obtained",
             "We kept the methodology and moved to a public dataset: Routine Blood Values (Huyut, Mendeley Data v1)",
             "Kept: ICU admission and death prediction, imbalance handling, feature selection, 5 classifiers, held-out test + bootstrap, data-type comparison",
             "Dropped: CT radiomics, clinical symptoms, ventilation outcome - not present in any public dataset we could find",
             "So our results are a lab-only study and are not directly comparable to the paper's numbers"])

sl = slide("Dataset", "Huyut, Routine Blood Values Datasets - Mendeley Data v1 (non-commercial / educational)")
tb = sl.shapes.add_table(3, 5, Inches(.6), Inches(1.9), Inches(12.1), Inches(1.8)).table
for j, h in enumerate(["Task", "Patients", "Positives", "Features", "Blocks"]): tb.cell(0, j).text = h
for i, row in enumerate([["ICU admission", f"{I['n']}", f"{I['positives']} ({I['positives']/I['n']:.1%})", "54", "biochemical 25, hematological 17, immunological 10, age, sex"],
                         ["In-hospital death", f"{D['n']}", f"{D['positives']} ({D['positives']/D['n']:.1%})", "37", "biochemical 12, hematological 15, immunological 10"]], 1):
    for j, v in enumerate(row): tb.cell(i, j).text = v
for j, wd in enumerate([2.3, 1.3, 1.9, 1.2, 5.4]): tb.columns[j].width = Inches(wd)
for r_ in tb.rows:
    for c in r_.cells:
        for p in c.text_frame.paragraphs:
            for run in p.runs: run.font.size = Pt(15)
bullets(sl, ["No patient ID: the three files per task are joined by row position (label order identical across files)",
             "ICU hematology/immunology files miss 1035 labels - labels taken from the biochemical file (they agree wherever both exist)"], y=4.9, size=18)

sl = slide("Data-quality problem we found (and handled)", "The most important finding for interpreting any result")
bullets(sl, ["Missing values were filled with constants. In the death files the constant differs by class: Ferritin = 395 for 92% of deaths; D-dimer = 1277 for 80% - the label leaks. A model on the raw data scores AUROC 1.00",
             "Fix: blank values covering >25% of a class in a column (or whose most common value differs by class); impossible zeros become missing",
             f"Even afterwards, a model that sees ONLY which values are missing reaches AUROC {I['diagnostics']['missingness_mask_only_auroc']:.2f} (ICU) / {D['diagnostics']['missingness_mask_only_auroc']:.2f} (death)",
             "So we impute with random draws from observed values (no constant for trees to detect) and treat results as likely optimistic"], x=.6, w=6.6, size=17)
img(sl, "data_quality.png", 7.4, 2.0, w=5.5)

sl = slide("Approach: pipeline and experiment design")
for k, (t, c) in enumerate([("Clean &\nalign", NAVY), ("Random-draw\nimpute, clip\n1-99%, scale", ACC), ("SMOTEENN\nor none", ACC), ("LASSO / PCA /\nFPR / none", ACC), ("LR, RF, SVM,\nMLP, LightGBM", ACC), ("Test once\n+ bootstrap", NAVY)]):
    box(sl, .6 + k * 2.08, 1.8, 1.9, 1.3, t, c, 14)
bullets(sl, ["Stratified 70/30 split; all 40 combinations per task scored by 5-fold CV on the 70% (resampling and selection refit inside every fold - no leakage)",
             "Best CV AUROC configuration refit on 70% and evaluated once on the 30% held-out test set",
             "Winner re-trained on each data block (biochemical / hematological / immunological / all) and compared with a plain LR baseline: paired bootstrap, 200 resamples, 95% CIs",
             "Threshold from out-of-fold scores (Youden); permutation importance for interpretation"], y=3.5, size=18)

sl = slide("Results on the held-out test set")
def row(d, k, nm): return [nm, f"{d['test'][k]['auroc']:.2f} ({d['test_ci'][k]['auroc'][0]:.2f}-{d['test_ci'][k]['auroc'][1]:.2f})", f"{d['test'][k]['auprc']:.2f}"]
rows = [["Data used", "ICU AUROC (95% CI)", "ICU AUPRC", "Death AUROC (95% CI)", "Death AUPRC"]]
for k, nm in [("All", "All lab features"), ("Biochemical", "Biochemical only"), ("Hematological", "Hematological only"), ("Immunological", "Immunological only"), ("Baseline LR", "Baseline LR")]:
    a, b = row(I, k, nm), row(D, k, nm); rows.append([nm, a[1], a[2], b[1], b[2]])
tb = sl.shapes.add_table(6, 5, Inches(.6), Inches(1.5), Inches(12.1), Inches(3)).table
for i, rw in enumerate(rows):
    for j, v in enumerate(rw):
        tb.cell(i, j).text = v
        for p in tb.cell(i, j).text_frame.paragraphs:
            for run in p.runs: run.font.size = Pt(17); run.font.bold = (i in (0, 1))
bullets(sl, [f"Best: ICU LightGBM + FPR selection (CV AUROC {I['cv_auroc']:.2f}); death LightGBM (CV AUROC {D['cv_auroc']:.2f}); no resampling in either",
             f"ICU: sensitivity {I['sensitivity']:.2f}, specificity {I['specificity']:.2f}, precision {I['precision']:.2f}   |   Death: {D['sensitivity']:.2f} / {D['specificity']:.2f} / {D['precision']:.2f}",
             "All blocks together beat every single block; for ICU, biochemistry carries nearly all the signal (hematology, immunology near chance)"], y=4.8, size=17)

sl = slide("ROC and precision-recall curves")
img(sl, "roc_pr_icu.png", .4, 1.4, w=6.3); img(sl, "roc_pr_death.png", 6.7, 1.4, w=6.3)
tx(sl, .6, 4.2, 12, 1, "Left: ICU admission. Right: in-hospital death. Dotted gray = baseline logistic regression.", 16, False, MUTE)
bullets(sl, ["Precision-recall matters more than ROC here: positives are only 14% (ICU) and 9% (death)"], y=5.0, size=18)

sl = slide("Robustness and interpretation")
img(sl, "bootstrap.png", .4, 1.4, w=6.2); img(sl, "importance.png", 6.6, 1.5, w=6.4)
bullets(sl, ["Bootstrap: all-lab model above each single block and the baseline in nearly all resamples on death; ICU margin over biochemistry alone is smaller",
             "ICU top features: ALT, calcium, age, AST, GGT; death: BASO, eGFR, RDW, troponin, procalcitonin"], y=5.4, size=16)

sl = slide("Differences from the paper")
bullets(sl, ["SMOTEENN did not help: best CV AUROC / AUPRC with it is lower than without it on both tasks (the paper used it in nearly all best models)",
             "FPR selection helped ICU (CV AUROC 0.91 vs 0.88 without); LASSO/PCA gave no gain on death",
             "Boosted trees won, as in the paper's ventilation task; the paper's gain from radiomics cannot be tested here",
             "Our 70/30 stratified split replaces the paper's date-based cohorts - no external validation cohort"])

sl = slide("Can we do better? Tested honestly", "Choices made with CV on the training split only; test set used once per model")
ICU, DTH = X2["icu"], X2["death"]
rows = [["Variant", "ICU CV AUROC", "Death CV AUROC", "Death CV AUPRC"]]
for k, nm in [("A baseline (current best)", "A  baseline"), ("B + ratio features", "B  + 6 clinical ratio features"), ("C + tuned LightGBM", "C  B + tuned LightGBM"),
              ("D C x 5 random imputations", "D  C averaged over 5 imputations"), ("E D + RF + ExtraTrees + LR", "E  D + RF + ExtraTrees + LR")]:
    rows.append([nm, f"{ICU['cv'][k]['auroc']:.3f}", f"{DTH['cv'][k]['auroc']:.3f}", f"{DTH['cv'][k]['auprc']:.3f}"])
tb = sl.shapes.add_table(6, 4, Inches(.6), Inches(1.7), Inches(8.2), Inches(2.9)).table
for i, rw in enumerate(rows):
    for j, v in enumerate(rw):
        tb.cell(i, j).text = v
        for p in tb.cell(i, j).text_frame.paragraphs:
            for run in p.runs: run.font.size = Pt(15); run.font.bold = (i == 0)
bullets(sl, [f"ICU: nothing beats the baseline - kept",
             f"Death: variant D wins; test AUPRC {DTH['test_base']['auprc']:.2f} to {DTH['test_new']['auprc']:.2f} (better in {DTH['prob_new_better_auprc']:.0%} of bootstraps), AUROC unchanged at {DTH['test_new']['auroc']:.2f}",
             f"Accuracy is misleading: always saying 'no' is {ICU['baseline_accuracy_always_negative']:.0%} / {DTH['baseline_accuracy_always_negative']:.0%} accurate",
             "We did not exploit the missing-value pattern (would give ~0.95-0.99 AUROC but reflects dataset artifacts)"], x=9.0, y=1.7, w=4.0, size=15)

sl = slide("Live demo and repository")
bullets(sl, ["Private GitHub repository, shared with faculty and TAs; README with setup and run steps",
             "python src/demo.py --task icu --n 6      scores random held-out patients: risk score, flag, true outcome",
             "python src/demo.py --task death --n 6",
             "python src/run.py reproduces everything in about two minutes (fixed seed 42); --quick is a smoke test"])
tx(sl, .6, 4.6, 12, 1.2, "src/data.py (load, align, clean)   |   src/pipeline.py (models)   |   src/run.py (experiments)   |   src/plots.py, demo.py, build_report.py", 16, False, MUTE)

sl = slide("Conclusions and limitations")
bullets(sl, [f"Lab-only models separate ICU-bound and fatal cases well on held-out data: AUROC {I['test']['All']['auroc']:.2f} (ICU), {D['test']['All']['auroc']:.2f} (death); fusing all blocks beat any single block",
             "Absolute numbers are probably optimistic: imputation artifacts, row-position join, no external cohort. Least-affected estimate: hematological-only death model AUROC "
             f"{D['test']['Hematological']['auroc']:.2f}",
             "Not clinically deployable as is - the dataset's artifacts would have to be fixed at collection time",
             "Future work: radiomics features, nested CV with hyperparameter search, calibration, external validation"], size=19)
P.save(root / "report" / "Slides.pptx"); print("saved")
